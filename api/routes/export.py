from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import FileResponse
from api.db import get_job
from fpdf import FPDF
from docx import Document
import os
import tempfile

router = APIRouter()

@router.get("/{job_id}/export")
async def export_job(job_id: str, format: str = Query(..., regex="^(pdf|word)$")):
    job = await get_job(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    if job.get("status") != "done":
        raise HTTPException(400, "Job is not yet complete")

    drafts = job.get("drafts", [])
    if not drafts:
        raise HTTPException(400, "No drafts found to export")

    if format == "pdf":
        return generate_pdf(job_id, drafts)
    else:
        return generate_word(job_id, drafts)

def generate_pdf(job_id, drafts):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Title
    pdf.set_font("Arial", 'B', 24)
    pdf.cell(200, 20, txt="Lumina AI Content Report", ln=True, align='C')
    pdf.ln(10)
    
    for draft in drafts:
        format_name = draft.get("format", "Unknown").upper()
        content = draft.get("content", "")
        
        pdf.set_font("Arial", 'B', 16)
        pdf.set_text_color(95, 94, 96) # Primary color
        pdf.cell(200, 10, txt=f"Format: {format_name}", ln=True)
        pdf.ln(5)
        
        pdf.set_font("Arial", size=11)
        pdf.set_text_color(0, 0, 0)
        # Using multi_cell for wrap
        pdf.multi_cell(0, 8, txt=content)
        pdf.ln(15)
        
    temp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(temp.name)
    return FileResponse(temp.name, filename=f"lumina_report_{job_id}.pdf")

def generate_word(job_id, drafts):
    doc = Document()
    doc.add_heading('Lumina AI Content Report', 0)
    
    for draft in drafts:
        format_name = draft.get("format", "Unknown").upper()
        content = draft.get("content", "")
        
        doc.add_heading(f'Format: {format_name}', level=1)
        doc.add_paragraph(content)
        doc.add_page_break()
        
    temp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(temp.name)
    return FileResponse(temp.name, filename=f"lumina_report_{job_id}.docx")
